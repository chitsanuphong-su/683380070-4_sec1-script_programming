import unittest
import os
import tempfile
import json
import docx
from PyPDF2 import PdfReader, PdfWriter

from src.utils import setup_logging, ensure_directory_exists, log_audit_entry
from src.config_parser import ConfigParser
from src.pdf_tasks import PDFTasks
from src.word_tasks import WordTasks
from src.document_agent import DocumentAgent


class TestAgenticDocumentAutomator(unittest.TestCase):

    def setUp(self):
        self.test_dir = tempfile.TemporaryDirectory()
        self.base_path = self.test_dir.name

    def tearDown(self):
        import logging
        for logger in [logging.getLogger("test_logger")]:
            for h in list(logger.handlers):
                h.close()
                logger.removeHandler(h)
        self.test_dir.cleanup()

    def test_utils_logging_and_directory(self):
        log_file = os.path.join(self.base_path, "test.log")
        logger = setup_logging("test_logger", filename=log_file)
        logger.info("Test log message")
        self.assertTrue(os.path.exists(log_file))

        sub_dir = os.path.join(self.base_path, "nested", "sub", "dir")
        ensure_directory_exists(sub_dir)
        self.assertTrue(os.path.isdir(sub_dir))

        audit_list = []
        log_audit_entry(audit_list, "SUCCESS", "All good", "test_task", {"key": "val"})
        self.assertEqual(len(audit_list), 1)
        self.assertEqual(audit_list[0]["status"], "SUCCESS")

    def test_config_parser_valid(self):
        cfg_path = os.path.join(self.base_path, "valid_cfg.json")
        valid_data = {
            "global_data": {"TITLE": "Demo"},
            "tasks": [
                {
                    "name": "Test Task",
                    "type": "generate_word_report",
                    "template_path": "tpl.docx",
                    "output_path": "out.docx"
                }
            ]
        }
        with open(cfg_path, "w", encoding="utf-8") as f:
            json.dump(valid_data, f)

        parser = ConfigParser(cfg_path)
        loaded = parser.load_config()
        self.assertEqual(loaded["global_data"]["TITLE"], "Demo")

    def test_config_parser_invalid_json(self):
        cfg_path = os.path.join(self.base_path, "invalid.json")
        with open(cfg_path, "w", encoding="utf-8") as f:
            f.write("{invalid json")
        parser = ConfigParser(cfg_path)
        with self.assertRaises(ValueError):
            parser.load_config()

    def test_config_parser_missing_fields(self):
        cfg_path = os.path.join(self.base_path, "missing.json")
        with open(cfg_path, "w", encoding="utf-8") as f:
            json.dump({"global_data": {}}, f)
        parser = ConfigParser(cfg_path)
        with self.assertRaises(ValueError):
            parser.load_config()

    def test_word_tasks_placeholder_and_table(self):
        tpl_path = os.path.join(self.base_path, "tpl.docx")
        out_path = os.path.join(self.base_path, "out.docx")

        doc = docx.Document()
        doc.add_paragraph("Hello {{NAME}}!")
        doc.add_paragraph("{{TABLE_PLACEHOLDER}}")
        doc.save(tpl_path)

        wt = WordTasks()
        success = wt.generate_report_from_template(
            tpl_path,
            out_path,
            {"{{NAME}}": "World"},
            table_data=[["Header1", "Header2"], ["A", "B"]]
        )
        self.assertTrue(success)
        self.assertTrue(os.path.exists(out_path))

        extracted = wt.extract_text_from_word(out_path)
        self.assertIn("Hello World!", extracted)
        self.assertIn("Header1", extracted)
        self.assertIn("A\tB", extracted)

    def test_pdf_tasks_watermark_merge_extract(self):
        from reportlab.pdfgen import canvas
        pdf1 = os.path.join(self.base_path, "p1.pdf")
        pdf2 = os.path.join(self.base_path, "p2.pdf")

        c1 = canvas.Canvas(pdf1)
        c1.drawString(100, 100, "Content Page One")
        c1.showPage()
        c1.save()

        c2 = canvas.Canvas(pdf2)
        c2.drawString(100, 100, "Content Page Two")
        c2.showPage()
        c2.save()

        pt = PDFTasks()
        # Merge
        merged = os.path.join(self.base_path, "merged.pdf")
        m_res = pt.merge_pdfs([pdf1, pdf2], merged)
        self.assertTrue(m_res)
        reader = PdfReader(merged)
        self.assertEqual(len(reader.pages), 2)

        # Watermark
        wm_out = os.path.join(self.base_path, "wm.pdf")
        w_res = pt.apply_watermark(pdf1, wm_out, "SECRET")
        self.assertTrue(w_res)
        self.assertTrue(os.path.exists(wm_out))

        # Extract
        txt = pt.extract_text_from_pdf(merged)
        self.assertIn("Content Page One", txt)
        self.assertIn("Content Page Two", txt)


if __name__ == "__main__":
    unittest.main()
