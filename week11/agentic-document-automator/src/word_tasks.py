# agentic-document-automator/src/word_tasks.py
import docx
from docx.oxml import OxmlElement
import os
try:
    from .utils import setup_logging
except ImportError:
    from utils import setup_logging

logger = setup_logging(__name__)


class WordTasks:
    """
    Handles advanced Word document operations for the Document Agent.
    """

    def __init__(self):
        logger.info("WordTasks initialized.")

    def _replace_placeholder_in_paragraph(self, paragraph, placeholder, value):
        """
        Replaces a single placeholder in a paragraph's text,
        safely handling placeholders that might span across multiple runs.
        """
        if placeholder not in paragraph.text:
            return False

        # First, try replacing within individual runs to preserve distinct styles if possible
        replaced_in_single_run = False
        for run in paragraph.runs:
            if placeholder in run.text:
                run.text = run.text.replace(placeholder, str(value))
                replaced_in_single_run = True

        if replaced_in_single_run:
            logger.debug(f"Replaced '{placeholder}' with '{value}' in single run.")
            return True

        # If placeholder spans multiple runs (common Word behavior):
        # We replace the text across runs while preserving the first run's formatting
        full_text = paragraph.text
        if placeholder in full_text:
            new_text = full_text.replace(placeholder, str(value))
            if paragraph.runs:
                # Keep first run's formatting and set all replaced text into it
                paragraph.runs[0].text = new_text
                # Clear remaining runs
                for run in paragraph.runs[1:]:
                    run.text = ""
            else:
                paragraph.text = new_text
            logger.debug(f"Replaced multi-run '{placeholder}' with '{value}' in paragraph.")
            return True

        return False

    def generate_report_from_template(self, template_path, output_path, data_placeholders, table_data=None):
        """
        Generates a Word document from a template, filling placeholders and optional table data.
        `data_placeholders` is a dictionary where keys are placeholders (e.g., '{{REPORT_TITLE}}')
        and values are the replacement strings.
        `table_data` is a list of lists, where the first sub-list is the header row.
        """
        if not os.path.exists(template_path):
            logger.error(f"Word template not found: {template_path}")
            return False

        try:
            document = docx.Document(template_path)

            # Replace placeholders in paragraphs
            for paragraph in document.paragraphs:
                for placeholder, value in data_placeholders.items():
                    self._replace_placeholder_in_paragraph(paragraph, placeholder, value)

            # Replace placeholders in table cells
            for table in document.tables:
                for row in table.rows:
                    for cell in row.cells:
                        for paragraph in cell.paragraphs:
                            for placeholder, value in data_placeholders.items():
                                self._replace_placeholder_in_paragraph(paragraph, placeholder, value)

            # Add dynamic table if provided
            if table_data:
                logger.info("Adding dynamic table to the document.")
                cols = len(table_data[0]) if table_data else 0

                # Find specific placeholder to insert table, or append to end
                table_anchor_found = False
                for paragraph in list(document.paragraphs):
                    if "{{TABLE_PLACEHOLDER}}" in paragraph.text:
                        table = document.add_table(rows=0, cols=cols)
                        table.style = 'Table Grid'

                        for r_idx, row_data in enumerate(table_data):
                            row = table.add_row()
                            for c_idx, cell_data in enumerate(row_data):
                                cell = row.cells[c_idx]
                                cell.text = str(cell_data)
                                if r_idx == 0:
                                    if cell.paragraphs and cell.paragraphs[0].runs:
                                        cell.paragraphs[0].runs[0].bold = True

                        # Move the table to right before the placeholder paragraph, then remove placeholder paragraph
                        p_elem = paragraph._element
                        p_elem.addprevious(table._tbl)
                        p_elem.getparent().remove(p_elem)

                        table_anchor_found = True
                        break

                if not table_anchor_found:
                    logger.warning("Table placeholder '{{TABLE_PLACEHOLDER}}' not found in template. Appending table to end.")
                    table = document.add_table(rows=0, cols=cols)
                    table.style = 'Table Grid'
                    for r_idx, row_data in enumerate(table_data):
                        row = table.add_row()
                        for c_idx, cell_data in enumerate(row_data):
                            cell = row.cells[c_idx]
                            cell.text = str(cell_data)
                            if r_idx == 0:
                                if cell.paragraphs and cell.paragraphs[0].runs:
                                    cell.paragraphs[0].runs[0].bold = True

            # Save the new document
            output_dir = os.path.dirname(output_path)
            if output_dir:
                os.makedirs(output_dir, exist_ok=True)
            document.save(output_path)
            logger.info(f"Report generated from template and saved to '{output_path}'.")
            return True
        except Exception as e:
            logger.error(f"Error generating report from template: {e}", exc_info=True)
            return False

    def extract_text_from_word(self, docx_path):
        """Extracts all text from a Word document."""
        if not os.path.exists(docx_path):
            logger.error(f"Word document not found: {docx_path}")
            return None

        try:
            document = docx.Document(docx_path)
            full_text = []
            for para in document.paragraphs:
                if para.text:
                    full_text.append(para.text)

            # Also extract text from tables
            for table in document.tables:
                for row in table.rows:
                    row_text = []
                    for cell in row.cells:
                        row_text.append(cell.text.strip())
                    full_text.append("\t".join(row_text))

            logger.info(f"Text extracted from Word document: {docx_path}")
            return "\n".join(full_text)
        except Exception as e:
            logger.error(f"Error extracting text from Word document '{docx_path}': {e}")
            return None
