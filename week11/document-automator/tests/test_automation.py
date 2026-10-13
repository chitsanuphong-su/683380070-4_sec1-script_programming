import os
import sys
import unittest
import tempfile
import shutil

# Add src to path
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), '..', 'src')))

from pdf_processor import PDFProcessor
from word_processor import WordProcessor
import docx
from PyPDF2 import PdfReader


class TestDocumentAutomation(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.base_dir = os.path.abspath(os.path.join(os.path.dirname(__file__), '..'))
        cls.docs_dir = os.path.join(cls.base_dir, 'documents')
        cls.input_article = os.path.join(cls.docs_dir, 'input_article.pdf')
        cls.dummy1 = os.path.join(cls.docs_dir, 'dummy_doc1.pdf')
        cls.dummy2 = os.path.join(cls.docs_dir, 'dummy_doc2.pdf')

    def setUp(self):
        self.test_dir = tempfile.mkdtemp()
        self.pdf_proc = PDFProcessor()
        self.word_proc = WordProcessor()

    def tearDown(self):
        shutil.rmtree(self.test_dir, ignore_errors=True)

    def test_extract_text_from_pdf(self):
        """Test extracting text from sample PDF."""
        self.assertTrue(os.path.exists(self.input_article))
        text = self.pdf_proc.extract_text_from_pdf(self.input_article)
        self.assertIsNotNone(text)
        self.assertIn("Introduction to Automation", text)
        self.assertIn("Benefits of Automation", text)
        self.assertIn("Conclusion", text)
        self.assertEqual(text.count("--- Page End ---"), 2)

    def test_extract_text_nonexistent_pdf(self):
        """Test extraction from non-existent file."""
        non_existent = os.path.join(self.test_dir, "non_existent.pdf")
        text = self.pdf_proc.extract_text_from_pdf(non_existent)
        self.assertIsNone(text)

    def test_merge_pdfs(self):
        """Test merging two valid PDFs."""
        out_pdf = os.path.join(self.test_dir, "merged_test.pdf")
        success = self.pdf_proc.merge_pdfs([self.dummy1, self.dummy2], out_pdf)
        self.assertTrue(success)
        self.assertTrue(os.path.exists(out_pdf))

        reader = PdfReader(out_pdf)
        self.assertEqual(len(reader.pages), 2)
        self.assertIn("Dummy Document 1", reader.pages[0].extract_text())
        self.assertIn("Dummy Document 2", reader.pages[1].extract_text())

    def test_merge_pdfs_empty_or_missing(self):
        """Test merge behavior with invalid or empty input list."""
        out_pdf = os.path.join(self.test_dir, "empty_merge.pdf")
        self.assertFalse(self.pdf_proc.merge_pdfs([], out_pdf))
        self.assertFalse(self.pdf_proc.merge_pdfs(["invalid_path.pdf"], out_pdf))

    def test_split_pdf(self):
        """Test splitting multi-page PDF."""
        split_out = os.path.join(self.test_dir, "splits")
        success = self.pdf_proc.split_pdf(self.input_article, output_dir=split_out, pages_per_split=1)
        self.assertTrue(success)

        files = sorted(os.listdir(split_out))
        self.assertEqual(len(files), 2)
        self.assertIn("input_article_part_1.pdf", files)
        self.assertIn("input_article_part_2.pdf", files)

        # Check content of each part
        r1 = PdfReader(os.path.join(split_out, "input_article_part_1.pdf"))
        self.assertEqual(len(r1.pages), 1)
        self.assertIn("Introduction to Automation", r1.pages[0].extract_text())

        r2 = PdfReader(os.path.join(split_out, "input_article_part_2.pdf"))
        self.assertEqual(len(r2.pages), 1)
        self.assertIn("Benefits of Automation", r2.pages[0].extract_text())

    def test_split_pdf_invalid_arguments(self):
        """Test split error handling with invalid pages_per_split and missing file."""
        self.assertFalse(self.pdf_proc.split_pdf("nonexistent.pdf"))
        self.assertFalse(self.pdf_proc.split_pdf(self.input_article, pages_per_split=0))
        self.assertFalse(self.pdf_proc.split_pdf(self.input_article, pages_per_split=-1))

    def test_word_processor_document_creation_and_formatting(self):
        """Test Word document creation with headings, formatting, and tables."""
        doc = self.word_proc.create_document()
        self.word_proc.add_heading(doc, "Test Title", level=1)
        self.word_proc.add_paragraph(doc, "Bold Text", bold=True, font_size=14, align="center")
        self.word_proc.add_paragraph(doc, "Italic Text", italic=True, align="right")

        table_data = [
            ["Col1", "Col2"],
            ["Val1", "Val2"]
        ]
        self.word_proc.add_table(doc, table_data, col_widths=[2.0, 2.0])

        out_docx = os.path.join(self.test_dir, "test_doc.docx")
        saved = self.word_proc.save_document(doc, out_docx)
        self.assertTrue(saved)
        self.assertTrue(os.path.exists(out_docx))

        # Reopen and inspect
        loaded_doc = docx.Document(out_docx)
        self.assertEqual(loaded_doc.paragraphs[0].text, "Test Title")
        self.assertEqual(loaded_doc.paragraphs[1].text, "Bold Text")
        self.assertTrue(loaded_doc.paragraphs[1].runs[0].bold)
        self.assertEqual(len(loaded_doc.tables), 1)
        self.assertEqual(loaded_doc.tables[0].cell(0, 0).text, "Col1")
        self.assertEqual(loaded_doc.tables[0].cell(1, 1).text, "Val2")

    def test_word_processor_add_picture_error_handling(self):
        """Test add_picture with nonexistent image."""
        doc = self.word_proc.create_document()
        result = self.word_proc.add_picture(doc, "non_existent_image.png")
        self.assertFalse(result)


if __name__ == '__main__':
    unittest.main()
